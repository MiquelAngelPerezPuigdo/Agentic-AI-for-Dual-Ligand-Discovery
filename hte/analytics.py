from __future__ import annotations

from collections import defaultdict
from pathlib import Path
import numpy as np

from .io import digest, finite, object_digest, read_csv, read_json, write_csv, write_json


def normalize_export(config, input_path, output):
    """Normalize vendor headers using the explicit HTE mapping, without guessing a detector."""
    path = Path(input_path)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[config["analytics"]["export_sheet"]]
        iterator = sheet.iter_rows(values_only=True)
        headers = next(iterator, None)
        if not headers or any(h is None for h in headers) or len(headers) != len(set(headers)):
            raise ValueError("Missing or duplicate spreadsheet export headers")
        raw = [dict(zip(headers, row)) for row in iterator if any(x is not None for x in row)]
    elif path.suffix.lower() == ".csv":
        raw = read_csv(path)
    else:
        raise ValueError("Supported integrated exports: CSV or XLSX. Convert legacy XLS with the vendor software.")
    mapping = config["analytics"]["column_mapping"]
    if not mapping:
        raise ValueError("HTE must supply the export column mapping")
    required = ["sample_id", "retention_time_min", "area", "channel", "method_id"]
    if any(key not in mapping for key in required):
        raise ValueError("Column mapping requires sample_id, retention_time_min, area, channel and method_id")
    normalized = []
    for r in raw:
        if any(mapping[key] not in r for key in required):
            raise ValueError("A mapped column is missing from the export")
        normalized.append({key: r[column] for key, column in mapping.items()})
    write_csv(output, normalized)


def fit_calibration(config, rows):
    settings = config["analytics"]
    mode = settings["calibration_mode"]
    if mode not in ("external", "internal_standard"):
        raise ValueError("Choose a validated external or internal-standard calibration")
    models = {}
    for analyte in ("substrate", "product"):
        points = [row for row in rows if row["analyte"] == analyte]
        if len(points) < 3:
            raise ValueError(f"At least three calibration measurements required for {analyte}")
        x, y, is_areas = [], [], []
        for row in points:
            if row["method_id"] != settings["method_id"] or row["channel"] != settings["channel"]:
                raise ValueError("Calibration method/channel mismatch")
            x.append(finite(row["concentration_uM"], "calibration concentration", 0))
            area = finite(row["area"], "calibration area", 0)
            if mode == "internal_standard":
                is_area = finite(row["internal_standard_area"], "calibration IS area", 1e-12)
                is_conc = finite(row["internal_standard_concentration_uM"], "IS concentration", 1e-12)
                if abs(is_conc-settings["internal_standard_final_uM"]) > 1e-6:
                    raise ValueError("Calibration and sample IS concentrations must match in the final LC vial")
                area /= is_area
                is_areas.append(is_area)
            y.append(area)
        if len(set(x)) < 3:
            raise ValueError(f"Three distinct calibration concentrations required for {analyte}")
        x, y = np.asarray(x), np.asarray(y)
        slope, intercept = np.linalg.lstsq(np.c_[x, np.ones(len(x))], y, rcond=None)[0]
        if not np.isfinite([slope, intercept]).all() or slope <= 0:
            raise ValueError("Invalid/nonpositive calibration slope")
        ss_total = float(np.sum((y-y.mean())**2))
        r2 = 1-float(np.sum((y-(slope*x+intercept))**2))/ss_total if ss_total else 0
        if r2 < settings["minimum_calibration_r2"]:
            raise ValueError(f"Calibration R2 for {analyte}: {r2:.5f} below acceptance threshold")
        models[analyte] = {"slope": float(slope), "intercept": float(intercept), "r2": r2,
                           "minimum_uM": float(x.min()), "maximum_uM": float(x.max()),
                           "mean_is_area": float(np.mean(is_areas)) if is_areas else None}
    return models


def validate_completion(ready_path, peaks_path, dosing_path, config):
    ready = read_json(ready_path)
    rows = read_csv(dosing_path)
    expected = [row["sample_id"] for row in rows]
    if ready.get("status") != "complete" or set(ready.get("sample_ids", [])) != set(expected):
        raise ValueError("LC completion manifest must explicitly cover the complete expected sample set")
    if len(ready["sample_ids"]) != len(expected):
        raise ValueError("Duplicate sample IDs in completion manifest")
    for field, path in [("peaks_sha256", peaks_path), ("dosing_sha256", dosing_path)]:
        if ready.get(field) != digest(path):
            raise ValueError(f"Completion manifest {field} does not match the supplied file")
    if ready.get("method_id") != config["analytics"]["method_id"]:
        raise ValueError("LC completion method differs from configured method")
    return ready


def complete_export(config, peaks_path, dosing_path, output):
    """Operator invokes this only after confirming every expected injection finished."""
    dosing = read_csv(dosing_path)
    expected = [r["sample_id"] for r in dosing]
    peaks = read_csv(peaks_path)
    present = {r["sample_id"] for r in peaks if r["channel"] == config["analytics"]["channel"]
               and r["method_id"] == config["analytics"]["method_id"]}
    if len(expected) != len(set(expected)) or present != set(expected):
        raise ValueError("Complete export requires every expected sample on the configured method/channel")
    write_json(output, {"status": "complete", "sample_ids": expected,
                        "method_id": config["analytics"]["method_id"],
                        "peaks_sha256": digest(peaks_path), "dosing_sha256": digest(dosing_path)})


def analyze(config, dosing, peaks, calibration, output):
    settings = config["analytics"]
    if not settings.get("method_id") or not settings.get("channel") or not settings.get("references"):
        raise ValueError("Configure LC method, detector channel and authentic-standard retention references")
    ids = [r["sample_id"] for r in dosing]
    if len(ids) != len(set(ids)):
        raise ValueError("Duplicate dosing sample IDs")
    groups = defaultdict(list)
    for row in peaks:
        if row["sample_id"] not in ids:
            raise ValueError(f"Unmapped LC sample: {row['sample_id']}")
        if row["channel"] != settings["channel"]:
            continue  # Different channels cannot be summed or substituted.
        if row["method_id"] != settings["method_id"]:
            raise ValueError("Sample LC method mismatch")
        finite(row["area"], "peak area", 0)
        finite(row["retention_time_min"], "retention time", 0)
        groups[row["sample_id"]].append(row)
    models = fit_calibration(config, calibration)
    loq = {name: finite(settings["quantification_limit_uM"][name], name+" LOQ", 0.000001)
           for name in ("substrate", "product")}
    if any(loq[name] > models[name]["maximum_uM"] for name in loq):
        raise ValueError("LOQ must lie within the calibration range")
    references = settings["references"]
    needed = ["substrate", "product"] + (["internal_standard"] if settings["calibration_mode"] == "internal_standard" else [])
    for name in needed:
        if name not in references:
            raise ValueError(f"Missing authentic-standard retention reference: {name}")
        finite(references[name]["retention_time_min"], name+" retention time", 0)
        finite(references[name]["window_min"], name+" RT window", 1e-6)
    # Ambiguous method windows must be fixed before samples are quantified.
    for i, a in enumerate(needed):
        for b in needed[i+1:]:
            if abs(references[a]["retention_time_min"]-references[b]["retention_time_min"]) <= references[a]["window_min"]+references[b]["window_min"]:
                raise ValueError(f"Overlapping retention windows: {a}, {b}")
    results = []
    for dose in dosing:
        flags, areas, censored = [], {}, {}
        sample_peaks = groups[dose["sample_id"]]
        for name in needed:
            ref = references[name]
            matches = [p for p in sample_peaks if abs(float(p["retention_time_min"])-ref["retention_time_min"]) <= ref["window_min"]]
            if len(matches) != 1:
                if not matches and name in loq and settings.get("missing_analyte_policy") == "censor":
                    censored[name] = loq[name]
                else:
                    flags.append(f"{name}_{'missing' if not matches else 'ambiguous'}")
                continue
            row = matches[0]
            if row.get("qc_flag", "").strip() not in ("", "ok"):
                flags.append(name+"_"+row["qc_flag"])
            areas[name] = float(row["area"])
        if settings["calibration_mode"] == "internal_standard" and areas.get("internal_standard", 0) <= 0:
            flags.append("invalid_internal_standard")
        elif settings["calibration_mode"] == "internal_standard":
            if any(abs(areas["internal_standard"]/model["mean_is_area"]-1) > settings["internal_standard_area_tolerance_fraction"] for model in models.values()):
                flags.append("internal_standard_area_outlier")
        concentrations = {}
        for name in ("substrate", "product"):
            if name not in areas:
                continue
            signal = areas[name]
            if settings["calibration_mode"] == "internal_standard":
                is_area = areas.get("internal_standard", 0)
                if is_area <= 0:
                    flags.append("invalid_internal_standard")
                    continue
                expected_is = models[name]["mean_is_area"]
                if abs(is_area/expected_is-1) > settings["internal_standard_area_tolerance_fraction"]:
                    flags.append("internal_standard_area_outlier")
                signal /= is_area
            model = models[name]
            conc = (signal-model["intercept"])/model["slope"]
            if conc < -settings["negative_concentration_tolerance_uM"]:
                flags.append(name+"_negative_concentration")
                continue
            conc = max(0, conc)
            if conc > model["maximum_uM"]+1e-6:
                flags.append(name+"_outside_calibration")
            if conc < loq[name]:
                censored[name] = loq[name]
            else:
                if conc < model["minimum_uM"]-1e-6:
                    flags.append(name+"_outside_calibration")
                concentrations[name] = conc
        extract_ul = float(dose["reaction_volume_ul"])+float(dose["workup_ul"])
        dilution = (float(dose["aliquot_ul"])+float(dose["diluent_ul"]))/float(dose["aliquot_ul"])
        amounts = {name: conc * dilution * extract_ul / 1e6 for name, conc in concentrations.items()}
        upper_amounts = {name: conc * dilution * extract_ul / 1e6 for name, conc in censored.items()}
        n0 = float(dose["substrate_umol"])
        y = amounts.get("product", 0)/n0*100 if n0 and "product" in amounts else None
        conversion = (1-amounts["substrate"]/n0)*100 if n0 and "substrate" in amounts else None
        if n0 and all(name in amounts for name in ("product", "substrate")):
            if amounts["product"]+amounts["substrate"] > n0*(1+settings["mass_balance_tolerance_fraction"]):
                flags.append("mass_balance_above_limit")
        if y is not None and y > 100+settings["yield_tolerance_percent"]:
            flags.append("yield_above_limit")
        if conversion is not None and conversion < -settings["yield_tolerance_percent"]:
            flags.append("negative_conversion")
        if dose["kind"] == "blank" and amounts.get("product", 0) > settings["blank_max_product_umol"]:
            flags.append("blank_contamination")
        # A valid nondetection is a bound, never a fabricated zero-yield measurement.
        results.append({"sample_id": dose["sample_id"], "round": dose["round"], "well": dose["well"],
                        "conditions_digest": object_digest(config["chemistry"]),
                        "kind": dose["kind"], "pair_id": dose["pair_id"], "ligand_a": dose["ligand_a"],
                        "ligand_b": dose["ligand_b"], "yield_percent": y, "conversion_percent": conversion,
                        "yield_upper_bound_percent": upper_amounts.get("product", 0)/n0*100 if n0 and "product" in upper_amounts else None,
                        "conversion_lower_bound_percent": (1-upper_amounts["substrate"]/n0)*100 if n0 and "substrate" in upper_amounts else None,
                        "censored_analytes": ";".join(sorted(censored)),
                        "product_umol": amounts.get("product"), "substrate_remaining_umol": amounts.get("substrate"),
                        "dilution_factor": dilution, "qc_pass": not flags, "qc_flags": ";".join(sorted(set(flags)))})
    rejected = sum(not r["qc_pass"] for r in results)
    control_failed = any(not r["qc_pass"] for r in results if r["kind"] == "blank")
    report = {"calibration": models, "samples": len(results), "rejected": rejected,
              "batch_qc_pass": not control_failed and rejected/len(results) <= settings["maximum_rejected_fraction"],
              "method_id": settings["method_id"], "channel": settings["channel"],
              "quantification": "Calibrated detector response; area percent is not yield."}
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    write_csv(output/"results.csv", results)
    write_json(output/"qc_report.json", report)
    return results, report


def observations(results):
    """Remove misleading invalid numeric values before sending feedback to the model."""
    output = []
    for row in results:
        valid = row["qc_pass"] is True or row["qc_pass"] == "True"
        value = {**row}
        if not valid:
            for key in ("yield_percent", "conversion_percent", "product_umol", "substrate_remaining_umol",
                        "yield_upper_bound_percent", "conversion_lower_bound_percent"):
                value[key] = None
        output.append(value)
    return output
