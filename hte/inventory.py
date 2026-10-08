from __future__ import annotations

from itertools import combinations
from rdkit import Chem
from rdkit.Chem import Descriptors

from .io import read_csv, write_csv


IDENTITY_FLAGS = {
    "1359986-21-2": "Name contains methoxy but supplied SMILES contains no oxygen. Confirm structure.",
    "29949-85-7": "Full name/SMILES indicate meta-chloro; common name says ortho. Confirm identity.",
}


def import_workbook(path, output):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["Reagents"]
    result = []
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
        if row_number == 1 or not row[0]:
            continue
        cas, name, smiles, common, cost, pack, vendor, url = row[:8]
        smiles = str(smiles).strip()
        mol = Chem.MolFromSmiles(smiles)
        if mol is None:
            raise ValueError(f"Unparseable structure at workbook row {row_number}")
        i = len(result)
        result.append({
            "ligand_id": f"L{i + 1:02d}", "cas": str(cas).strip(),
            "name": str(common or name).strip(), "full_name": str(name).strip(),
            "smiles": smiles, "molecular_weight_g_mol": round(Descriptors.MolWt(mol), 4),
            "phosphorus_atoms": sum(a.GetAtomicNum() == 15 for a in mol.GetAtoms()),
            "source_slot": 1 if i < 24 else 4,
            "source_well": f"{'ABCD'[i % 24 % 4]}{i % 24 // 4 + 1}",
            "identity_confirmed": "false" if str(cas).strip() in IDENTITY_FLAGS else "true",
            "identity_note": IDENTITY_FLAGS.get(str(cas).strip(), ""),
            "source_workbook_row": row_number,
        })
    if len(result) != 31:
        raise ValueError(f"Expected 31 ligands; found {len(result)}")
    write_csv(output, result)
    return result


def load_inventory(path, confirmed=False):
    rows = read_csv(path)
    seen, locations = set(), set()
    for row in rows:
        key = row["ligand_id"]
        if key in seen or not key:
            raise ValueError(f"Duplicate or empty ligand ID: {key}")
        seen.add(key)
        slot, well = int(row["source_slot"]), row["source_well"]
        if slot not in (1, 4) or well not in [f"{r}{c}" for c in range(1, 7) for r in "ABCD"]:
            raise ValueError(f"Invalid ligand source address: {key}")
        if (slot, well) in locations:
            raise ValueError("Duplicate ligand source location")
        locations.add((slot, well))
        if Chem.MolFromSmiles(row["smiles"]) is None:
            raise ValueError(f"Invalid ligand SMILES: {key}")
        if confirmed and row.get("identity_confirmed") != "true":
            raise ValueError(f"Unresolved identity {key}: {row.get('identity_note', '')}")
    return {row["ligand_id"]: row for row in rows}


def pair_id(a, b):
    if a == b:
        raise ValueError("Dual-ligand candidates must contain distinct ligands")
    return "__".join(sorted((a, b)))


def candidates(inventory):
    return [{"pair_id": pair_id(a, b), "ligand_a": a, "ligand_b": b}
            for a, b in combinations(sorted(inventory), 2)]
